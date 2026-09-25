#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
jobscz_stats.py — z trackeru vygeneruje deterministicky statisticky souhrn
pro posledni zmerene obdobi: reports/stats_RRRR-MM.json (+ reports/latest.json).

Souhrn cte naplanovana uloha v Claude Cowork a pise z nej mesicni report.
Vsechna cisla se pocitaji zde (presne a opakovatelne), Claude je jen interpretuje.

Nazvy pozic a firem jsou data z verejneho webu (neduveryhodna) - pouze se
kopiruji jako text, nic se z nich neinterpretuje ani nespousti.

Bezi v GitHub Actions po kazdem dennim behu. Obsah nezavisi na case spusteni,
takze se soubor meni (a commituje) jen kdyz se zmeni data v trackeru.
"""
import os, re, json, datetime, unicodedata, collections
from openpyxl import load_workbook

TRACKER = "AI_jobs_tracker_jobscz.xlsx"
OUT_DIR = "reports"
SUMMARY = "Souhrn (trend)"
SNAP = "Inzer\u00e1ty (snapshot)"
FIRST_ROW = 5
PERIOD_COL = 11
REPO = "https://github.com/kozubektp/jobscz-ai-tracker"
MAX_NEW_A = 40
MAX_GONE_A = 20

def last_day(year, month):
    nxt = datetime.date(year + (month == 12), month % 12 + 1, 1)
    return nxt - datetime.timedelta(days=1)

def months_between(p1, p2):
    y1, m1 = map(int, p1.split("-")); y2, m2 = map(int, p2.split("-"))
    return (y2 - y1) * 12 + (m2 - m1)

def safe_text(s, limit=150):
    s = "".join(ch for ch in str(s or "") if unicodedata.category(ch)[0] != "C")
    s = re.sub(r"\s+", " ", s).strip()
    return s[:limit]

def num(v):
    return v if isinstance(v, (int, float)) else 0

def read_series(ws):
    rows, r, blanks = {}, FIRST_ROW, 0
    while r < 5000 and blanks < 5:
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            blanks += 1; r += 1; continue
        blanks = 0
        date = str(a)[:10]
        period = str(ws.cell(row=r, column=PERIOD_COL).value or date[:7])
        total, matched, promoted, A = (num(ws.cell(row=r, column=c).value) for c in (2, 3, 4, 6))
        if total and matched:
            ai_rel = matched - promoted
            rows[period] = {
                "period": period, "date": date, "total": total, "matched": matched,
                "promoted": promoted, "ai_rel": ai_rel, "A": A, "B": ai_rel - A,
                "pct_ai_of_total": round(ai_rel / total * 100, 2),
                "pct_A_of_total": round(A / total * 100, 3),
                "pct_A_of_ai": round(A / ai_rel * 100, 1) if ai_rel else 0,
            }
        r += 1
    return [rows[k] for k in sorted(rows)]

def read_listings(ws2):
    by_date = collections.defaultdict(list)
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        by_date[str(row[0])[:10]].append({
            "id": str(row[1]), "cat": row[2], "title": safe_text(row[3]),
            "company": safe_text(row[4], 80), "promoted": row[5] == "ano"})
    return by_date

def norm_title(t):
    t = t.lower()
    t = re.sub(r"\((m/[fž](/[dx])?|ž/m|f/m)\)", "", t)
    return re.sub(r"\s+", " ", t).strip()

def top_companies(listings, cat=None, n=10):
    c = collections.Counter(x["company"] for x in listings
                            if x["company"] and not x["promoted"] and (cat is None or x["cat"] == cat))
    return [{"company": k, "count": v} for k, v in sorted(c.items(), key=lambda kv: (-kv[1], kv[0]))[:n]]

def delta(prev, last, key):
    a, b = prev[key], last[key]
    return {"prev": a, "last": b, "abs": b - a, "pct": round((b - a) / a * 100, 1) if a else None}

def main():
    wb = load_workbook(TRACKER, read_only=True)
    series = read_series(wb[SUMMARY])
    listings = read_listings(wb[SNAP])
    if not series:
        print("[stats] tracker nema data, nic negeneruji"); return 0
    last = series[-1]
    prev = series[-2] if len(series) > 1 else None
    y, m = map(int, last["period"].split("-"))
    L = [x for x in listings.get(last["date"], [])]
    LA = {x["id"]: x for x in L if x["cat"] == "A" and not x["promoted"]}

    out = {
        "schema": 1,
        "period": last["period"],
        "snapshot_date": last["date"],
        "final": last["date"] >= last_day(y, m).isoformat(),
        "source": REPO,
        "tracker_url": f"{REPO}/blob/main/{TRACKER}",
        "series": series,
        "latest": last,
        "previous": prev,
        "top_A_companies_latest": top_companies(L, "A"),
        "top_all_companies_latest": top_companies(L, None),
        "A_dedup_latest": len({(norm_title(x["title"]), x["company"]) for x in LA.values()}),
        "notes": [],
    }

    if prev:
        P = [x for x in listings.get(prev["date"], [])]
        PA = {x["id"]: x for x in P if x["cat"] == "A" and not x["promoted"]}
        Lid, Pid = {x["id"] for x in L}, {x["id"] for x in P}
        out["gap_months"] = months_between(prev["period"], last["period"]) - 1
        out["deltas"] = {k: delta(prev, last, k) for k in ("total", "ai_rel", "A", "B")}
        out["share_change_pp"] = {k: round(last[k] - prev[k], 3)
                                  for k in ("pct_ai_of_total", "pct_A_of_total", "pct_A_of_ai")}
        out["churn"] = {
            "all": {"survived": len(Lid & Pid), "new": len(Lid - Pid), "gone": len(Pid - Lid)},
            "A": {"survived": len(set(LA) & set(PA)), "new": len(set(LA) - set(PA)),
                  "gone": len(set(PA) - set(LA))},
            "listing_data_available": bool(L) and bool(P),
        }
        new_A = sorted((LA[i] for i in set(LA) - set(PA)), key=lambda x: (x["company"] == "", x["company"], x["title"]))
        gone_A = sorted((PA[i] for i in set(PA) - set(LA)), key=lambda x: (x["company"] == "", x["company"], x["title"]))
        out["new_A_roles"] = [f'{x["title"]} | {x["company"]}' for x in new_A[:MAX_NEW_A]]
        out["new_A_roles_truncated"] = max(0, len(new_A) - MAX_NEW_A)
        out["gone_A_roles_sample"] = [f'{x["title"]} | {x["company"]}' for x in gone_A[:MAX_GONE_A]]
        out["top_A_companies_previous"] = top_companies(P, "A")
        if out["gap_months"] > 0:
            out["notes"].append(f'Mezi obdobimi {prev["period"]} a {last["period"]} chybi '
                                f'{out["gap_months"]} mesic(e) mereni.')
    if not out["final"]:
        out["notes"].append("Mereni neni z posledniho dne mesice (predbezne / rucni).")
    if out["A_dedup_latest"] < last["A"]:
        out["notes"].append(f'Po ocisteni o duplicitni inzeraty (stejna role CZ/EN) je A = '
                            f'{out["A_dedup_latest"]} misto {last["A"]}.')

    os.makedirs(OUT_DIR, exist_ok=True)
    text = json.dumps(out, ensure_ascii=False, indent=1, sort_keys=True)
    for name in (f"stats_{last['period']}.json", "latest.json"):
        with open(os.path.join(OUT_DIR, name), "w", encoding="utf-8") as f:
            f.write(text + "\n")
    print(f"[stats] {OUT_DIR}/stats_{last['period']}.json (final={out['final']})")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
