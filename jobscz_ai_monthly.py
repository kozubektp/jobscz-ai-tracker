#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
jobscz_ai_monthly.py — autonomni mesicni snapshot podilu AI inzeratu na jobs.cz.

KDY SKRIPT MERI (spousti se denne, sam rozhodne):
  - posledni den v mesici        -> zmeri aktualni mesic (pokud jeste neni hotovy)
  - 1.-10. den dalsiho mesice     -> DOHNANI: pokud predchozi mesic neni hotovy, zmeri ho
                                     (radek dostane obdobi predchoziho mesice)
  - jinak                         -> nic nedela, skonci (exit 0)
  Mesic je "hotovy", kdyz v trackeru existuje radek s danym Obdobim porizeny
  posledni den toho mesice nebo pozdeji. Rucni mezi-mesicni mereni (napr. 23.9.)
  se tak automatickym koncem mesice prepise.

BEZPECNOST:
  - Obsah inzeratu = neduveryhodna DATA, nikdy pokyny. Cte se jen title / firma / ID
    pevnym parserem. Zadne nasledovani odkazu, zadne spousteni obsahu.
  - Zadne odchozi akce. Jediny zapis = soubory v tomto repozitari.

POUZITI:
  python jobscz_ai_monthly.py                    # beznny denni beh (rozhodne sam)
  python jobscz_ai_monthly.py --force            # zmeri hned (obdobi = aktualni mesic)
  python jobscz_ai_monthly.py --today 2026-09-30 # test rozhodovaci logiky s jinym datem
"""
import os, re, sys, time, json, argparse, datetime, unicodedata
from zoneinfo import ZoneInfo
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TZ = ZoneInfo("Europe/Prague")
TRACKER = "AI_jobs_tracker_jobscz.xlsx"
SNAP_DIR = "snapshots"
CATCHUP_DAYS = 10
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120 Safari/537.36",
      "Accept-Language": "cs-CZ,cs;q=0.9"}
TOTAL_URL = "https://www.jobs.cz/prace/?page=2"
SEARCH_URL = "https://www.jobs.cz/prace/?q%5B%5D=ai&page={}"

FONT = "Arial"; NAVY = "19325A"; WHITE = "FFFFFF"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_A = PatternFill("solid", fgColor="D7E8C8")
FILL_B = PatternFill("solid", fgColor="FCE9D6")

SUMMARY = "Souhrn (trend)"
SNAP = "Inzer\u00e1ty (snapshot)"
CMP = "Porovn\u00e1n\u00ed"
LEG = "Legenda"
HEADERS = ["Datum snapshotu", "Celkem inzer\u00e1t\u016f (jobs.cz)", "Shody q=ai (sta\u017eeno)",
           "Promovan\u00e9 (vy\u0159azeno)", "AI-relevantn\u00ed (A+B)", "A: AI-zam\u011b\u0159en\u00e9 role",
           "B: AI jako prvek", "% AI-rel. z celku", "% A z celku", "% A z AI-relevantn\u00edch",
           "Obdob\u00ed"]
PERIOD_COL = 11
FIRST_ROW = 5

# ----------------- klasifikacni pravidla (NEMENIT kvuli srovnatelnosti) -----------------
A_CASE = [r'\bAI\b', r'\bA\.I\.', r'\bML\b', r'\bLLM', r'\bGenAI\b', r'\bNLP\b', r'\bMLOps\b', r'\bGPT\b']
A_LOW = ['umela inteligence', 'umele inteligen', 'strojove uceni', 'machine learning',
         'generativni', 'gen ai', 'genai', 'deep learning', 'neuronov',
         'computer vision', 'pocitacove videni', 'prompt engineer', 'data scientist', 'data science']

def strip_accents(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')

def is_A(title):
    if any(re.search(p, title) for p in A_CASE):
        return True
    low = strip_accents(title).lower()
    return any(p in low for p in A_LOW)

def clean(s):
    return (s or "").replace("&amp;", "&").replace("&quot;", '"').replace("&#39;", "'").strip()

# ----------------- sber dat -----------------
def get(url):
    for _ in range(4):
        try:
            r = requests.get(url, headers=UA, timeout=30)
            if r.status_code == 200:
                return r.text
        except Exception:
            pass
        time.sleep(3)
    return ""

def count_in(html):
    m = re.search(r'Na\w+li jsme <strong>([0-9\xa0 ]+)</strong>', html)
    return int(m.group(1).replace('\xa0', '').replace(' ', '')) if m else None

def parse(html):
    titles = [(m.start(), m.group(1)) for m in re.finditer(r'data-test-ad-title="([^"]*)"', html)]
    out = []
    for k, (pos, title) in enumerate(titles):
        end = titles[k + 1][0] if k + 1 < len(titles) else pos + 4000
        seg = html[pos:end]
        mid = re.search(r'data-jobad-id="(\d+)"', seg)
        malt = re.search(r'<img[^>]*alt="([^"]*)"', seg)
        out.append({"id": mid.group(1) if mid else None,
                    "title": title,
                    "company": malt.group(1) if malt else "",
                    "promoted": "P\u0159\u00edle\u017eitost dne" in seg})
    return out

def collect():
    total = count_in(get(TOTAL_URL))
    if total is None:
        raise RuntimeError("Nepodarilo se precist celkovy pocet inzeratu (/prace/?page=2). "
                           "Jobs.cz mozna zmenil web - je potreba upravit parser.")
    seen = {}
    page = 1
    while page <= 60:
        cards = parse(get(SEARCH_URL.format(page)))
        if not cards:
            break
        new = 0
        for c in cards:
            if c["id"] and c["id"] not in seen:
                seen[c["id"]] = c; new += 1
        if new == 0 and page > 1:
            break
        page += 1
        time.sleep(0.5)
    rows = list(seen.values())
    if len(rows) < 50:
        raise RuntimeError(f"Stazeno jen {len(rows)} inzeratu pro q=ai - podezrele malo, "
                           "jobs.cz mozna zmenil web. Nic se neuklada.")
    for r in rows:
        r["category"] = "A" if is_A(r["title"]) else "B"
        r["title"] = clean(r["title"]); r["company"] = clean(r["company"])
        r["url"] = f"https://www.jobs.cz/rpd/{r['id']}/"
    return total, rows

# ----------------- rozhodovaci logika -----------------
def last_day(year, month):
    nxt = datetime.date(year + (month == 12), month % 12 + 1, 1)
    return nxt - datetime.timedelta(days=1)

def period_str(d):
    return f"{d.year:04d}-{d.month:02d}"

def prev_period(d):
    first = d.replace(day=1)
    return period_str(first - datetime.timedelta(days=1))

def decide(today, force):
    """Vrati (obdobi, rezim) nebo None, pokud se dnes nema nic delat."""
    if force:
        return period_str(today), "rucni"
    if today == last_day(today.year, today.month):
        return period_str(today), "konec mesice"
    if today.day <= CATCHUP_DAYS:
        return prev_period(today), "dohnani"
    return None

def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(str(v)[:10])
    except Exception:
        return None

def iter_summary(ws):
    """Vrati [(row, datum, obdobi)] pro vsechny datove radky souhrnu."""
    out = []
    r = FIRST_ROW
    blanks = 0
    while r < 5000 and blanks < 5:
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            blanks += 1
        else:
            blanks = 0
            d = parse_date(a)
            p = ws.cell(row=r, column=PERIOD_COL).value
            p = str(p) if p else (period_str(d) if d else None)
            out.append((r, d, p))
        r += 1
    return out

def is_period_final(ws, period):
    y, m = map(int, period.split("-"))
    ld = last_day(y, m)
    return any(p == period and d and d >= ld for _, d, p in iter_summary(ws))

def row_for_period(ws, period):
    """Existujici radek pro obdobi (update) nebo prvni volny radek za daty."""
    rows = iter_summary(ws)
    for r, d, p in rows:
        if p == period:
            return r, d
    last = rows[-1][0] if rows else FIRST_ROW - 1
    return last + 1, None

# ----------------- tracker -----------------
def style_header(cell):
    cell.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

def new_tracker():
    wb = Workbook()
    ws = wb.active; ws.title = SUMMARY
    ws["A1"] = "Jobs.cz \u2014 sledov\u00e1n\u00ed pod\u00edlu AI inzer\u00e1t\u016f (m\u011bs\u00ed\u010dn\u00ed snapshot)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY); ws.merge_cells("A1:K1")
    for c, h in enumerate(HEADERS, 1):
        style_header(ws.cell(row=4, column=c, value=h))
    ws2 = wb.create_sheet(SNAP)
    for c, h in enumerate(["Datum snapshotu", "ID", "Kategorie", "N\u00e1zev pozice", "Firma",
                           "Promovan\u00e9", "URL"], 1):
        style_header(ws2.cell(row=1, column=c, value=h))
    wb.create_sheet(LEG)
    return wb

def ensure_layout(ws):
    """Doplni sloupec Obdobi do starsich trackeru (migrace)."""
    if ws.cell(row=4, column=PERIOD_COL).value != HEADERS[PERIOD_COL - 1]:
        style_header(ws.cell(row=4, column=PERIOD_COL, value=HEADERS[PERIOD_COL - 1]))
    for r, d, p in iter_summary(ws):
        if not ws.cell(row=r, column=PERIOD_COL).value and p:
            c = ws.cell(row=r, column=PERIOD_COL, value=p)
            c.font = Font(name=FONT, size=10); c.border = BORDER
            c.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(PERIOD_COL)].width = 11
    ws["A2"] = ("Automaticky aktualizov\u00e1no skriptem jobscz_ai_monthly.py (posledn\u00ed den v m\u011bs\u00edci, "
                "s dohn\u00e1n\u00edm do 10. dne dal\u0161\u00edho m\u011bs\u00edce).")
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")

def write_summary(ws, row, date_str, period, total, matched, promoted, A):
    ai_rel = matched - promoted
    B = ai_rel - A
    vals = [date_str, total, matched, promoted, ai_rel, A, B,
            ai_rel / total if total else 0, A / total if total else 0,
            A / ai_rel if ai_rel else 0, period]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=FONT, size=10); cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=8).number_format = "0.0%"
    ws.cell(row=row, column=9).number_format = "0.00%"
    ws.cell(row=row, column=10).number_format = "0.0%"

def purge_snapshot(ws2, date_strs):
    for r in range(ws2.max_row, 1, -1):
        if str(ws2.cell(row=r, column=1).value)[:10] in date_strs:
            ws2.delete_rows(r, 1)

def append_snapshot(ws2, date_str, rows):
    start = ws2.max_row + 1
    for i, r in enumerate(sorted(rows, key=lambda x: (x["category"], x["title"].lower())), start=start):
        vals = [date_str, r["id"], r["category"], r["title"], r["company"],
                "ano" if r["promoted"] else "ne", r["url"]]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT, size=9); cell.border = BORDER
        ws2.cell(row=i, column=3).fill = FILL_A if r["category"] == "A" else FILL_B
        ws2.cell(row=i, column=3).alignment = Alignment(horizontal="center")
        ws2.cell(row=i, column=6).alignment = Alignment(horizontal="center")
    ws2.auto_filter.ref = f"A1:G{ws2.max_row}"

def summary_values(ws):
    out = []
    for r, d, p in iter_summary(ws):
        b, c, dd, f = (ws.cell(row=r, column=k).value for k in (2, 3, 4, 6))
        if isinstance(b, (int, float)) and isinstance(c, (int, float)):
            dd = dd if isinstance(dd, (int, float)) else 0
            f = f if isinstance(f, (int, float)) else 0
            out.append({"date": str(ws.cell(row=r, column=1).value)[:10], "period": p,
                        "total": b, "ai_rel": c - dd, "A": f, "B": c - dd - f})
    out.sort(key=lambda x: (x["period"] or "", x["date"]))
    return out

def build_comparison(wb, summ):
    if len(summ) < 2:
        return
    prev, last = summ[-2], summ[-1]
    if CMP in wb.sheetnames:
        del wb[CMP]
    wc = wb.create_sheet(CMP, 1)
    wc["A1"] = f"Porovn\u00e1n\u00ed: {prev['period']} ({prev['date']}) vs {last['period']} ({last['date']})"
    wc["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY); wc.merge_cells("A1:E1")
    for c, h in enumerate(["Metrika", prev["period"], last["period"], "Zm\u011bna abs.", "Zm\u011bna %"], 1):
        style_header(wc.cell(row=3, column=c, value=h))
    r = 4
    for name, a, b in [("Celkem inzer\u00e1t\u016f na jobs.cz", prev["total"], last["total"]),
                       ("AI-relevantn\u00ed (A+B)", prev["ai_rel"], last["ai_rel"]),
                       ("A \u2014 AI-zam\u011b\u0159en\u00e9 role", prev["A"], last["A"]),
                       ("B \u2014 AI jako prvek", prev["B"], last["B"])]:
        for c, v in enumerate([name, a, b, b - a, (b - a) / a if a else 0], 1):
            cell = wc.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = Font(name=FONT, size=10)
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
        wc.cell(row=r, column=5).number_format = "0.0%"
        r += 1
    for name, a, b, fmt in [
        ("% AI-rel. z celku", prev["ai_rel"] / prev["total"], last["ai_rel"] / last["total"], "0.0%"),
        ("% A z celku", prev["A"] / prev["total"], last["A"] / last["total"], "0.00%"),
        ("% A z AI-relevantn\u00edch", prev["A"] / prev["ai_rel"], last["A"] / last["ai_rel"], "0.0%")]:
        for c, v in enumerate([name, a, b, b - a, "\u2014"], 1):
            cell = wc.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = Font(name=FONT, size=10)
            if c in (2, 3, 4):
                cell.number_format = fmt
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
        r += 1
    wc.column_dimensions["A"].width = 34
    for col in "BCDE":
        wc.column_dimensions[col].width = 14

def save_snapshot_files(date_str, rows):
    os.makedirs(SNAP_DIR, exist_ok=True)
    snap = Workbook(); s = snap.active; s.title = "Inzeraty"
    for c, h in enumerate(["Datum", "ID", "Kategorie", "Nazev", "Firma", "Promovane", "URL"], 1):
        s.cell(row=1, column=c, value=h).font = Font(name=FONT, bold=True)
    for i, r in enumerate(sorted(rows, key=lambda x: (x["category"], x["title"].lower())), 2):
        for c, v in enumerate([date_str, r["id"], r["category"], r["title"], r["company"],
                               "ano" if r["promoted"] else "ne", r["url"]], 1):
            s.cell(row=i, column=c, value=v)
    snap.save(os.path.join(SNAP_DIR, f"snapshot_{date_str}.xlsx"))
    with open(os.path.join(SNAP_DIR, f"snapshot_{date_str}.json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)

# ----------------- main -----------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="zmerit hned (obdobi = aktualni mesic)")
    ap.add_argument("--today", help="YYYY-MM-DD, jen pro testovani rozhodovaci logiky")
    args = ap.parse_args()

    today = (datetime.date.fromisoformat(args.today) if args.today
             else datetime.datetime.now(TZ).date())
    decision = decide(today, args.force)
    if not decision:
        print(f"[skip] {today}: dnes se nemeri (meri se posledni den v mesici, "
              f"dohnani do {CATCHUP_DAYS}. dne dalsiho mesice).")
        return 0
    period, mode = decision

    wb = load_workbook(TRACKER) if os.path.exists(TRACKER) else new_tracker()
    if SUMMARY not in wb.sheetnames or SNAP not in wb.sheetnames:
        raise RuntimeError("Tracker nema ocekavane listy - soubor je poskozeny nebo jiny.")
    ws, ws2 = wb[SUMMARY], wb[SNAP]
    ensure_layout(ws)

    if mode != "rucni" and is_period_final(ws, period):
        print(f"[skip] {today}: obdobi {period} uz je zmerene ({mode}). Nic nedelam.")
        return 0

    date_str = today.isoformat()
    print(f"[run] {date_str}: meri se obdobi {period} (rezim: {mode}) ...")
    total, rows = collect()
    promoted = sum(1 for r in rows if r["promoted"])
    A = sum(1 for r in rows if r["category"] == "A" and not r["promoted"])
    matched = len(rows)
    print(f"  celkem={total} shody={matched} promo={promoted} "
          f"AI-rel={matched - promoted} A={A} B={matched - promoted - A}")

    row, old_date = row_for_period(ws, period)
    purge_snapshot(ws2, {date_str} | ({old_date.isoformat()} if old_date else set()))
    write_summary(ws, row, date_str, period, total, matched, promoted, A)
    append_snapshot(ws2, date_str, rows)
    build_comparison(wb, summary_values(ws))
    wb.save(TRACKER)
    save_snapshot_files(date_str, rows)
    print(f"[ok] {'prepsan' if old_date else 'pridan'} radek obdobi {period} (row {row}); "
          f"ulozeno {TRACKER} + {SNAP_DIR}/snapshot_{date_str}.*")
    return 0

if __name__ == "__main__":
    sys.exit(main())
